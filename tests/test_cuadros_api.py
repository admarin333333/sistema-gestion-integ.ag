# Prueba de API: cuadros de composición de las notas (tanda 3)
import json
import urllib.request
import urllib.error

BASE = "http://127.0.0.1:8010"
ok = 0
fallos = []


def req(metodo, ruta, datos=None, token=None):
    r = urllib.request.Request(BASE + ruta, method=metodo)
    r.add_header("Content-Type", "application/json")
    if token:
        r.add_header("Authorization", "Bearer " + token)
    cuerpo = json.dumps(datos).encode() if datos is not None else None
    try:
        with urllib.request.urlopen(r, cuerpo, timeout=30) as resp:
            return resp.status, json.loads(resp.read().decode())
    except urllib.error.HTTPError as e:
        return e.code, json.loads(e.read().decode() or "{}")


def check(nombre, cond, detalle=""):
    global ok
    if cond:
        ok += 1
        print(f"  OK  {nombre}")
    else:
        fallos.append(nombre)
        print(f"FALLA {nombre} -> {detalle}")


# Login
st, tok = req("POST", "/api/auth/login", {"usuario": "admin", "password": "admin123"})
token = tok.get("access_token")
check("login", st == 200 and token, f"status {st}")

# GET ejercicio 1 --------------------------------------------------------
st, ej = req("GET", "/api/balance-rt54/ejercicios/1", token=token)
check("GET ejercicio", st == 200, f"status {st}")

notas = {n["clave"]: n for n in ej["notas"]}
c21 = notas["2.1"]["cuadro"]
check("2.1 tiene cuadro", c21 is not None)
check(
    "2.1 columnas",
    [(c["letra"], c["grupo"], c["etiqueta"]) for c in c21["columnas"]]
    == [("C", None, "Actual"), ("D", None, "Comparativo")],
    str(c21["columnas"]),
)
filas21 = {f["fila"]: f for f in c21["filas"]}
check("2.1 filas 20-24", sorted(filas21) == [20, 21, 22, 23, 24], str(sorted(filas21)))
check("2.1 rotulo fila 20", filas21[20]["rotulo"] == "Efectivo Fondo fijo", filas21[20]["rotulo"])
check("2.1 rotulo fila 24", filas21[24]["rotulo"] == "Total", filas21[24]["rotulo"])
check("2.1 fila 24 es calc", filas21[24]["celdas"]["C"]["calc"] is True)
check("2.1 celda sin cargar = null", filas21[20]["celdas"]["C"]["valor"] is None, str(filas21[20]["celdas"]["C"]))
check("2.1 total vacío = 0", filas21[24]["celdas"]["C"]["valor"] == 0)

c23 = notas["2.3"]["cuadro"]
check(
    "2.3 columnas 4",
    [(c["letra"], c["grupo"], c["etiqueta"]) for c in c23["columnas"]]
    == [
        ("C", "Corriente", "Actual"),
        ("D", "Corriente", "Comparativo"),
        ("E", "No corriente", "Actual"),
        ("F", "No corriente", "Comparativo"),
    ],
    str(c23["columnas"]),
)
filas23 = {f["fila"]: f for f in c23["filas"]}
check("2.3 tipos", filas23[49]["tipo"] == "subtotal" and filas23[50]["tipo"] == "prevision" and filas23[51]["tipo"] == "total")
check("2.12 tiene cuadro", notas["2.12"]["cuadro"] is not None)
check("3.4 tiene cuadro", notas["3.4"]["cuadro"] is not None)
check("2.6 sin cuadro", notas["2.6"]["cuadro"] is None)
check("2.20 sin cuadro (renglones en blanco del modelo)", notas["2.20"]["cuadro"] is None)
check("2.22 sin cuadro (renglones en blanco del modelo)", notas["2.22"]["cuadro"] is None)
check("1.1 sin cuadro", notas["1.1"]["cuadro"] is None)
check("nota 1.1 texto guardado sigue ahí", notas["1.1"]["texto"] == "Nota 1.1 editada desde la pantalla.", notas["1.1"]["texto"][:40])

# Rotulo con espacio normal (en el modelo hay \xa0)
rot110 = {f["fila"]: f for f in notas["2.9"]["cuadro"]["filas"]}[110]["rotulo"]
check("2.9 rotulo sin \xa0", "\xa0" not in rot110 and "\n" not in rot110, repr(rot110))

# Claves inválidas --------------------------------------------------------
st, r = req("PUT", "/api/balance-rt54/ejercicios/1", {
    "cabecera": ej["cabecera"], "valores": [],
    "celdas_nota": [{"clave": "2.1:99:C", "valor": 5}],
}, token)
check("clave fila inexistente -> 400", st == 400, str(st) + str(r))
st, r = req("PUT", "/api/balance-rt54/ejercicios/1", {
    "cabecera": ej["cabecera"], "valores": [],
    "celdas_nota": [{"clave": "2.1:24:C", "valor": 5}],
}, token)
check("celda Total (no cargable) -> 400", st == 400, str(st))
st, r = req("PUT", "/api/balance-rt54/ejercicios/1", {
    "cabecera": ej["cabecera"], "valores": [],
    "celdas_nota": [{"clave": "2.20:188:C", "valor": 5}],
}, token)
check("nota sin cuadro -> 400", st == 400, str(st))
st, r = req("PUT", "/api/balance-rt54/ejercicios/1", {
    "cabecera": ej["cabecera"], "valores": [],
    "celdas_nota": [{"clave": "2.1:20:Z", "valor": 5}],
}, token)
check("columna inexistente -> 400", st == 400, str(st))

# Guardar celdas ----------------------------------------------------------
celdas = [
    {"clave": "2.1:20:C", "valor": 100.5},
    {"clave": "2.1:21:C", "valor": 30},
    {"clave": "2.1:23:C", "valor": 500},
    {"clave": "2.3:44:C", "valor": 10},
    {"clave": "2.3:45:C", "valor": 20},
    {"clave": "2.3:46:C", "valor": 30},
    {"clave": "2.3:47:C", "valor": 40},
    {"clave": "2.3:48:C", "valor": 50},
    {"clave": "2.3:50:C", "valor": -15},
    {"clave": "2.3:44:D", "valor": 1.11},
    {"clave": "2.3:45:D", "valor": 2.22},
    {"clave": "2.3:44:E", "valor": 7},
    {"clave": "3.4:276:C", "valor": 1000},
    {"clave": "3.4:278:C", "valor": 400},
]
st, r = req("PUT", "/api/balance-rt54/ejercicios/1", {
    "cabecera": ej["cabecera"], "valores": [], "celdas_nota": celdas,
}, token)
check("PUT con celdas -> 200", st == 200, str(st) + str(r)[:200])

notas2 = {n["clave"]: n for n in r["notas"]}
f21 = {f["fila"]: f for f in notas2["2.1"]["cuadro"]["filas"]}
check("2.1 C20 persiste", f21[20]["celdas"]["C"]["valor"] == 100.5, str(f21[20]["celdas"]["C"]))
check("2.1 C22 sigue null", f21[22]["celdas"]["C"]["valor"] is None)
check("2.1 total C = 630.5", f21[24]["celdas"]["C"]["valor"] == 630.5, str(f21[24]["celdas"]["C"]["valor"]))
check("2.1 total D = 0", f21[24]["celdas"]["D"]["valor"] == 0)

f23 = {f["fila"]: f for f in notas2["2.3"]["cuadro"]["filas"]}
check("2.3 Subtotal C = 150", f23[49]["celdas"]["C"]["valor"] == 150, str(f23[49]["celdas"]["C"]))
check("2.3 Subtotal D = 3.33", f23[49]["celdas"]["D"]["valor"] == 3.33, str(f23[49]["celdas"]["D"]))
check("2.3 Total C = 135 (150 - 15)", f23[51]["celdas"]["C"]["valor"] == 135, str(f23[51]["celdas"]["C"]))
check("2.3 Subtotal E = 7", f23[49]["celdas"]["E"]["valor"] == 7, str(f23[49]["celdas"]["E"]))
check("2.3 Total E = 7 (sin previsión cargada)", f23[51]["celdas"]["E"]["valor"] == 7, str(f23[51]["celdas"]["E"]))
check("2.3 Total F = 0", f23[51]["celdas"]["F"]["valor"] == 0)

f34 = {f["fila"]: f for f in notas2["3.4"]["cuadro"]["filas"]}
check("3.4 Total Otros ingresos = 1000", f34[277]["celdas"]["C"]["valor"] == 1000, str(f34[277]["celdas"]["C"]))
check("3.4 Total Otros egresos = 400 (bloque reiniciado)", f34[279]["celdas"]["C"]["valor"] == 400, str(f34[279]["celdas"]["C"]))
check("2.4 sin tocar -> todo null", all(
    c["valor"] is None
    for f in notas2["2.4"]["cuadro"]["filas"] if f["tipo"] in ("dato", "prevision")
    for c in f["celdas"].values()
))

# Export -------------------------------------------------------------------
import urllib.request as u
r2 = u.Request(BASE + "/api/balance-rt54/ejercicios/1/export.xlsx", headers={"Authorization": "Bearer " + token})
with u.urlopen(r2, timeout=60) as resp:
    datos = resp.status, resp.read()
st, data = datos
check("export 200", st == 200, str(st))
path = r"C:\Users\admar\AppData\Local\Temp\opencode\eecc\cuadros_export.xlsx"
import pathlib
pathlib.Path(path).parent.mkdir(parents=True, exist_ok=True)
open(path, "wb").write(data)

import openpyxl  # noqa: E402
wb = openpyxl.load_workbook(path, data_only=False)
ws = wb["Notas"]
check("Excel C20 = 100.5", ws["C20"].value == 100.5, str(ws["C20"].value))
check("Excel C22 sin cargar = vacío", ws["C22"].value is None, str(ws["C22"].value))
check("Excel C24 (Total) = 630.5", ws["C24"].value == 630.5, str(ws["C24"].value))
check("Excel D24 (Total) = 0", ws["D24"].value == 0, str(ws["D24"].value))
check("Excel rotulo B20 intacto", ws["B20"].value == "Efectivo Fondo fijo", str(ws["B20"].value))
check("Excel C49 (Subtotal) = 150", ws["C49"].value == 150, str(ws["C49"].value))
check("Excel C50 (Previsión) = -15", ws["C50"].value == -15, str(ws["C50"].value))
check("Excel C51 (Total) = 135", ws["C51"].value == 135, str(ws["C51"].value))
check("Excel D49 = 3.33", ws["D49"].value == 3.33, str(ws["D49"].value))
check("Excel E49 = 7", ws["E49"].value == 7, str(ws["E49"].value))
check("Excel C277 = 1000", ws["C277"].value == 1000, str(ws["C277"].value))
check("Excel C279 = 400", ws["C279"].value == 400, str(ws["C279"].value))
check("Excel 2.4 intacta (C57 vacío)", ws["C57"].value is None, str(ws["C57"].value))
check("Excel 2.4 Total intacto (C59 vacío)", ws["C59"].value is None, str(ws["C59"].value))
check("Excel 2.15 intacta (C143 vacío)", ws["C143"].value is None, str(ws["C143"].value))
check("Excel nota 1.1 texto", "Nota 1.1 editada desde la pantalla." in str(ws["B7"].value or ""), str(ws["B7"].value)[:80])

# Vaciar: lista sin celdas -----------------------------------------------
st, r = req("PUT", "/api/balance-rt54/ejercicios/1", {
    "cabecera": ej["cabecera"], "valores": [], "celdas_nota": [],
}, token)
check("PUT lista vacía -> 200", st == 200, str(st))
notas3 = {n["clave"]: n for n in r["notas"]}
f21b = {f["fila"]: f for f in notas3["2.1"]["cuadro"]["filas"]}
check("vacío: C20 null", f21b[20]["celdas"]["C"]["valor"] is None, str(f21b[21]["celdas"]["C"]))
check("vacío: total C = 0", f21b[24]["celdas"]["C"]["valor"] == 0)

r2 = u.Request(BASE + "/api/balance-rt54/ejercicios/1/export.xlsx", headers={"Authorization": "Bearer " + token})
with u.urlopen(r2, timeout=60) as resp:
    data2 = resp.read()
path2 = r"C:\Users\admar\AppData\Local\Temp\opencode\eecc\cuadros_export_vacio.xlsx"
open(path2, "wb").write(data2)
wb2 = openpyxl.load_workbook(path2, data_only=False)
ws2 = wb2["Notas"]
check("export vacío: C20 intacto", ws2["C20"].value is None, str(ws2["C20"].value))
check("export vacío: C24 intacto", ws2["C24"].value is None, str(ws2["C24"].value))

print()
print(f"RESULTADO: {ok} OK, {len(fallos)} fallos")
for f in fallos:
    print("  -", f)
