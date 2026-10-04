# Prueba de los campos nuevos de la CARATULA (RPC, controlante y capital)
import io
import json
import os
import sys
import urllib.request
import urllib.error

import openpyxl
import warnings

warnings.filterwarnings("ignore")
# La dirección de la API sale de `GC_BASE_URL` para poder correr estas pruebas
# contra la base de PRUEBAS y no contra la del estudio. Si la variable no
# está, usa 8010 como antes: no cambia cómo se corren.

BASE = os.environ.get("GC_BASE_URL", "http://127.0.0.1:8010")

# La suite necesita el ejercicio RT54 número 1, que está en la base del estudio y
# no en la de pruebas. Si no está, se avisa y se omite la suite en vez de
# reventar con un `KeyError: 'cabecera'` que no dice nada.
# Ver `datos_de_prueba.py`.
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from datos_de_prueba import chequear_o_salir  # noqa: E402
# El backend, calculado desde donde esta este archivo: asi el proyecto se
# puede mover de carpeta sin romper los tests.
_RUTA_BACKEND = os.path.normpath(
    os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "backend")
)
PLANTILLA = os.path.join(
    _RUTA_BACKEND,
    "plantillas",
    # Ojo: el nombre va SIN la barra inicial. Con `os.path.join`, un componente
    # que empieza con "\" se toma como ruta absoluta y descarta el resto.
    "MODELO-EECC-Entes-Con-Fines-de-Lucro-RT54-publicacion_new.xlsx",
)
ok = 0
fallos = []


def req(metodo, ruta, datos=None, token=None):
    r = urllib.request.Request(BASE + ruta, method=metodo)
    r.add_header("Content-Type", "application/json")
    if token:
        r.add_header("Authorization", "Bearer " + token)
    cuerpo = json.dumps(datos).encode() if datos is not None else None
    try:
        with urllib.request.urlopen(r, cuerpo, timeout=30) as resp:
            return resp.status, json.loads(resp.read().decode())
    except urllib.error.HTTPError as e:
        return e.code, json.loads(e.read().decode() or "{}")


def descargar(ruta, token):
    r = urllib.request.Request(BASE + ruta)
    r.add_header("Authorization", "Bearer " + token)
    with urllib.request.urlopen(r, timeout=60) as resp:
        return resp.read()


def check(nombre, cond, detalle=""):
    global ok
    if cond:
        ok += 1
        print(f"  OK  {nombre}")
    else:
        fallos.append(nombre)
        print(f"FALLA {nombre} -> {detalle}")


st, tok = req("POST", "/api/auth/login", {"usuario": "admin", "password": "admin123"})
token = tok.get("access_token")
check("login", st == 200 and token, str(st))

chequear_o_salir("caratula", token, BASE)

# Valores originales para restaurar al final
st, original = req("GET", "/api/balance-rt54/ejercicios/1", token=token)
check("lee ejercicio", st == 200, str(st))
cab0 = dict(original["cabecera"])

# ---------------------------------------------------------------- guardar ---
nuevos = {
    "fecha_inscripcion_rpc": "2020-03-15",
    "fecha_estatuto": "2019-12-01",
    "fecha_modificacion_estatuto": "2021-06-30",
    "fecha_vencimiento_entidad": "2049-12-31",
    "matricula_rpc": "IMA-123456",
    "identificacion_rpc": "Acta 45/2020",
    "duracion_entidad": "30",          # llega como texto desde el input
    "unidad_medida": "Pesos",
    "controlante_denominacion": "Holdco SA",
    "controlante_domicilio": "Av. Siempreviva 742",
    "controlante_actividad": "Inversiones",
    "controlante_participacion": "85%",
    "controlante_votos": "90%",
    "entes_nota": "5",
    "cap_circ_cantidad": "1000",
    "cap_circ_tipo": "Ordinarias",
    "cap_circ_votos": "1",
    "cap_circ_suscripto": "100000",
    "cap_circ_integrado": "",           # vacio tiene que guardar None
    "cap_cart_cantidad": 50,
    "cap_cart_tipo": "Preferentes",
    "cap_cart_votos": 2,
    "cap_cart_suscripto": "5000,50",    # coma decimal del teclado
    "cap_cart_integrado": 5000,
}
cab_nueva = dict(cab0)
cab_nueva.update(nuevos)
st, r = req(
    "PUT",
    "/api/balance-rt54/ejercicios/1",
    {"cabecera": cab_nueva, "valores": [], "celdas_nota": []},
    token,
)
check("guarda campos nuevos", st == 200, str(st) + " " + str(r))

st, guardado = req("GET", "/api/balance-rt54/ejercicios/1", token=token)
cab = guardado["cabecera"]


def mismo(esperado, obtenido):
    """Texto igual; números comparados como números; "" = quedó vacío."""
    if esperado == "":
        return obtenido is None
    if isinstance(esperado, (int, float)):
        return obtenido is not None and abs(float(esperado) - float(obtenido)) < 0.001
    try:
        return abs(float(str(esperado).replace(",", ".")) - float(obtenido)) < 0.001
    except (ValueError, TypeError):
        return str(obtenido) == str(esperado)


for campo, valor_ in nuevos.items():
    check(f"persiste {campo}", mismo(valor_, cab.get(campo)),
          f"{campo}={cab.get(campo)!r} esperaba {valor_!r}")

# ----------------------------------------------------------------- excel ----
plantilla = openpyxl.load_workbook(PLANTILLA, data_only=False)["CARATULA"]
datos = descargar("/api/balance-rt54/ejercicios/1/export.xlsx", token)
wb = openpyxl.load_workbook(io.BytesIO(datos), data_only=False)
check("abre el Excel", "CARATULA" in wb.sheetnames, str(wb.sheetnames))
ws = wb["CARATULA"]


def valor(celda):
    return str(ws[celda].value)


# Fechas: el marcador "..." del modelo se reemplaza por dd/mm/aaaa
f_insc = plantilla["C20"].value.split("\u2026")[0]
check("C20 fecha inscripcion", valor("C20") == f_insc + "15/03/2020", valor("C20"))
f_est = "Del Estatuto o Contrato Social de fecha: 01/12/2019"
check("C22 fecha estatuto",
      valor("C22").startswith(f_est) and valor("C22").endswith("."),
      valor("C22"))
check(
    "G22 matricula",
    valor("G22") == plantilla["G22"].value.replace("\u2026", "IMA-123456"),
    valor("G22"),
)
f_mod = plantilla["C23"].value.split("\u2026")[0]
check("C23 ultima modificacion", valor("C23") == f_mod + "30/06/2021", valor("C23"))
f_id = plantilla["C24"].value.split("\u2026")[0]
check("C24 identificacion RPC", valor("C24") == f_id + "Acta 45/2020", valor("C24"))
check(
    "G26 duracion",
    valor("G26") == plantilla["G26"].value.replace("\u2026", "30"),
    valor("G26"),
)
check("G27 fecha vencimiento", valor("G27") == "31/12/2049", valor("G27"))

# Sin marcador: etiqueta del modelo + ": " + valor
etiquetas = {
    "C28": "Pesos",
    "C29": "Holdco SA",
    "C30": "Av. Siempreviva 742",
    "C31": "Inversiones",
    "C32": "85%",
    "C33": "90%",
    "C35": "5",
}
for celda, v in etiquetas.items():
    etiqueta = str(plantilla[celda].value).rstrip()
    if not etiqueta.endswith(":"):
        etiqueta += ":"
    check(f"{celda} etiqueta+valor", valor(celda) == etiqueta + " " + v, valor(celda))

# Composicion del capital: dos renglones con numeros de verdad
fila45 = {"C": 1000, "D": "Ordinarias", "E": 1, "G": 100000.0,
          "H": plantilla["H45"].value}
fila52 = {"C": 50, "D": "Preferentes", "E": 2, "G": 5000.5, "H": 5000.0}
for col, esperado in fila45.items():
    check(f"{col}45 circulacion", ws[f"{col}45"].value == esperado,
          f"{ws[f'{col}45'].value!r} != {esperado!r}")
for col, esperado in fila52.items():
    check(f"{col}52 cartera", ws[f"{col}52"].value == esperado,
          f"{ws[f'{col}52'].value!r} != {esperado!r}")
check("C45 es numero", isinstance(ws["C45"].value, (int, float)), type(ws["C45"].value))
check("G52 es numero", isinstance(ws["G52"].value, (int, float)), type(ws["G52"].value))

# ------------------------------------------------- vacio = modelo intacto ---
cab_vacia = dict(cab)
for campo in nuevos:
    cab_vacia[campo] = None if campo != "duracion_entidad" else None
cab_vacia.update({
    "nombre": cab["nombre"], "fecha_inicio": cab["fecha_inicio"],
    "fecha_fin": cab["fecha_fin"], "entidad": cab["entidad"],
    "cuit": cab["cuit"], "domicilio": cab["domicilio"],
    "actividad": cab["actividad"], "actividad_secundaria": cab["actividad_secundaria"],
})
st, r = req(
    "PUT",
    "/api/balance-rt54/ejercicios/1",
    {"cabecera": cab_vacia, "valores": [], "celdas_nota": []},
    token,
)
check("guarda vacios", st == 200, str(st) + " " + str(r))
datos2 = descargar("/api/balance-rt54/ejercicios/1/export.xlsx", token)
ws2 = openpyxl.load_workbook(io.BytesIO(datos2), data_only=False)["CARATULA"]
check("C20 intacta si vacia", "\u2026" in str(ws2["C20"].value), str(ws2["C20"].value))
check("G26 intacta si vacia", "\u2026" in str(ws2["G26"].value), str(ws2["G26"].value))
check("C28 intacta si vacia", str(ws2["C28"].value) == str(plantilla["C28"].value),
      str(ws2["C28"].value))
check("C45 intacta si vacia", str(ws2["C45"].value) == str(plantilla["C45"].value),
      str(ws2["C45"].value))
check("D52 intacta si vacia", str(ws2["D52"].value) == str(plantilla["D52"].value),
      str(ws2["D52"].value))

# --------------------------------------------------------------- restore ----
st, r = req(
    "PUT",
    "/api/balance-rt54/ejercicios/1",
    {"cabecera": cab0, "valores": [], "celdas_nota": []},
    token,
)
check("restaura caratula original", st == 200, str(st))

print(f"\n{ok} OK / {len(fallos)} fallos")
if fallos:
    raise SystemExit(1)
