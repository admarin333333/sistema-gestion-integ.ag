# Prueba de índices de moneda homogénea + solapa Moneda homogénea
import os
import sys
import io
import json
import urllib.request
import urllib.error
from decimal import Decimal, ROUND_HALF_UP

import openpyxl
import warnings

warnings.filterwarnings("ignore")
# La dirección de la API sale de `GC_BASE_URL` para poder correr estas pruebas
# contra la base de PRUEBAS y no contra la del estudio. Si la variable no
# está, usa 8010 como antes: no cambia cómo se corren.

BASE = os.environ.get("GC_BASE_URL", "http://127.0.0.1:8010")

# Esta suite necesita los índices de moneda homogénea (los 404 valores del IPC de
# AFIP). Están en la base del estudio, cargados a mano: no hay ningún archivo del
# proyecto que los tenga y la API no los genera. Sin ellos, la prueba reventaba
# con `IndexError: list index out of range` al pedir el primero de la lista.
# Ver `datos_de_prueba.py`.
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from datos_de_prueba import chequear_o_salir  # noqa: E402
ok = 0
fallos = []


def req(metodo, ruta, datos=None, token=None):
    r = urllib.request.Request(BASE + ruta, method=metodo)
    r.add_header("Content-Type", "application/json")
    if token:
        r.add_header("Authorization", "Bearer " + token)
    cuerpo = json.dumps(datos).encode() if datos is not None else None
    try:
        with urllib.request.urlopen(r, cuerpo, timeout=30) as resp:
            texto = resp.read().decode()
            return resp.status, json.loads(texto) if texto else None
    except urllib.error.HTTPError as e:
        texto = e.read().decode() or "{}"
        try:
            return e.code, json.loads(texto)
        except json.JSONDecodeError:
            return e.code, {"detalle": texto}


def descargar(ruta, token):
    r = urllib.request.Request(BASE + ruta)
    r.add_header("Authorization", "Bearer " + token)
    with urllib.request.urlopen(r, timeout=60) as resp:
        return resp.read()


def check(nombre, cond, detalle=""):
    global ok
    if cond:
        ok += 1
        print(f"  OK  {nombre}")
    else:
        fallos.append(nombre)
        print(f"FALLA {nombre} -> {detalle}")


def cuatro(v):
    return Decimal(str(v)).quantize(Decimal("0.0001"), rounding=ROUND_HALF_UP)


def pesos(v):
    return Decimal(str(v)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


st, tok = req("POST", "/api/auth/login", {"usuario": "admin", "password": "admin123"})
token = tok.get("access_token")
check("login", st == 200 and token, str(st))

chequear_o_salir("moneda", token, BASE)

# ------------------------------------------------------------ índices -------
st, indices = req("GET", "/api/indices-moneda", token=token)
check("lista índices", st == 200 and len(indices) >= 400, f"{st} {len(indices or [])}")
check("más nuevo primero", indices[0]["fecha"] >= indices[-1]["fecha"],
      f"{indices[0]['fecha']} vs {indices[-1]['fecha']}")
por_mes = {i["fecha"]: i for i in indices}

st, r = req("POST", "/api/indices-moneda", {"fecha": "2027-01-15", "indice": 13000.5}, token)
check("alta de índice", st == 201 and r["fecha"] == "2027-01-01", f"{st} {r}")
id_prueba = r["id"]
st, r2 = req("POST", "/api/indices-moneda", {"fecha": "2027-01-01", "indice": 13100.25}, token)
check("corrige el mismo mes", st == 201 and r2["id"] == id_prueba
      and abs(r2["indice"] - 13100.25) < 0.0001, f"{st} {r2}")
st, lista2 = req("GET", "/api/indices-moneda", token=token)
check("no duplica el mes", sum(1 for i in lista2 if i["fecha"] == "2027-01-01") == 1,
      str(len(lista2)))
st, r = req("POST", "/api/indices-moneda", {"fecha": "2027-02-01", "indice": -5}, token)
check("rechaza índice negativo", st == 422, str(st))
st, _ = req("DELETE", f"/api/indices-moneda/{id_prueba}", token=token)
check("borra índice de prueba", st == 204, str(st))


def payload_cliente(c):
    return {
        "tipo_persona": c["tipo_persona"],
        "nombre": c["nombre"],
        "apellido": c["apellido"],
        "cuit": c["cuit"],
        "dni": c["dni"],
        "email": c["email"],
        "cod_area": c["cod_area"],
        "telefono": c["telefono"],
        "calle": c["calle"],
        "numero_calle": c["numero_calle"],
        "localidad": c["localidad"],
        "codigo_postal": c["codigo_postal"],
        "provincia": c["provincia"],
        "actividad_economica": c["actividad_economica"],
        "tipo_actividad": c["tipo_actividad"],
        "condicion_iva": c["condicion_iva"],
        "alicuota_iva_id": (c.get("alicuota_iva") or {}).get("id")
        or c.get("alicuota_iva_id"),
        "observaciones": c["observaciones"],
        "fecha_cierre_ejercicio": c.get("fecha_cierre_ejercicio"),
        "servicios": [s["id"] for s in c.get("servicios") or []],
    }


# ------------------------------------------------- ejercicio1 (sin ancla) ---
# El cliente13 es RI sin alícuota: su ficha no se puede guardar (así venía),
# así que su ejercicio usa el camino "sin fecha de cierre" = cierre +1 año.
st, ej = req("GET", "/api/balance-rt54/ejercicios/1", token=token)
check("lee ejercicio1", st == 200, str(st))
fecha_fin = ej["cabecera"]["fecha_fin"]
st, m = req("GET", "/api/balance-rt54/ejercicios/1/moneda-homogenea", token=token)
check("moneda homogénea responde (cierre +1 año)", st == 200, str(st) + " " + str(m)[:200])
if st == 200:
    anio_esperado = str(int(fecha_fin[:4]) + 1)
    check("sin ancla: próximo cierre = mismo día +1 año",
          m["proxima_fecha_cierre"] == f"{anio_esperado}{fecha_fin[4:]}",
          f"{m['proxima_fecha_cierre']} vs {anio_esperado}{fecha_fin[4:]}")
    idx_cierre = por_mes.get(fecha_fin[:7] + "-01", {}).get("indice")
    idx_prox = por_mes.get(m["proxima_fecha_cierre"][:7] + "-01", {}).get("indice")
    check("índice del cierre", idx_cierre is not None
          and abs(m["indice_anterior"]["valor"] - float(cuatro(idx_cierre))) < 1e-9,
          f"{m['indice_anterior']['valor']} vs {idx_cierre}")
    check("índice del próximo cierre", idx_prox is not None
          and abs(m["indice_nuevo"]["valor"] - float(cuatro(idx_prox))) < 1e-9,
          f"{m['indice_nuevo']['valor']} vs {idx_prox}")
    coef_esperado = (cuatro(idx_prox) / cuatro(idx_cierre)).quantize(
        Decimal("0.0001"), rounding=ROUND_HALF_UP
    )
    check("coeficiente con4 decimales",
          abs(Decimal(str(m["coeficiente"])) - coef_esperado) < Decimal("0.0000001"),
          f"{m['coeficiente']} vs {coef_esperado}")

    fila_rubro = next((f for f in m["esp"]["activo"] if f["tipo"] == "rubro"), None)
    check("hay rubros en el activo", fila_rubro is not None, str(m["esp"]["activo"][:2]))
    if fila_rubro:
        original_actual = fila_rubro.get("actual", 0)
        original_anterior = fila_rubro.get("anterior", 0)
        st, resp = req(
            "PUT",
            "/api/balance-rt54/ejercicios/1",
            {
                "cabecera": ej["cabecera"],
                "valores": [{
                    "seccion": "esp",
                    "clave": fila_rubro["clave"],
                    "valor_actual": 1000,
                    "valor_anterior": 0,
                }],
                "celdas_nota": [],
            },
            token,
        )
        check("guarda importe de prueba", st == 200, str(st) + str(resp)[:150])
        st, m2 = req("GET", "/api/balance-rt54/ejercicios/1/moneda-homogenea", token=token)
        fila = next(f for f in m2["esp"]["activo"] if f["clave"] == fila_rubro["clave"])
        esperado = pesos(Decimal("1000") * Decimal(str(m2["coeficiente"])))
        check("1000 x coeficiente = actualizado",
              abs(Decimal(str(fila["actualizado"])) - esperado) < Decimal("0.0001"),
              f"{fila['actualizado']} vs {esperado}")
        malas = []
        for seccion in ("activo", "pasivo"):
            for f in m2["esp"][seccion]:
                if f["tipo"] in ("rubro", "total") and f.get("actualizado") is not None:
                    esp = pesos(Decimal(str(f["actual"])) * Decimal(str(m2["coeficiente"])))
                    if abs(Decimal(str(f["actualizado"])) - esp) > Decimal("0.0001"):
                        malas.append(f["clave"])
        for f in m2["er"]:
            if f.get("actualizado") is not None:
                esp = pesos(Decimal(str(f["actual"])) * Decimal(str(m2["coeficiente"])))
                if abs(Decimal(str(f["actualizado"])) - esp) > Decimal("0.0001"):
                    malas.append("er:" + f["clave"])
        check("todos los renglones actualizados al coeficiente", not malas, str(malas))

        datos = descargar("/api/balance-rt54/ejercicios/1/moneda-homogenea.xlsx", token)
        wb = openpyxl.load_workbook(io.BytesIO(datos), data_only=True)
        ws = wb.active
        check("abre el Excel de moneda", ws.title == "Moneda homogénea", ws.title)
        a4 = str(ws["A4"].value)
        check("A4 dice índice del cierre", "Índice fecha de cierre anterior" in a4
              and fecha_fin[8:10] + "/" + fecha_fin[5:7] + "/" + fecha_fin[:4] in a4, a4)
        check("A4 con4 decimales", "," in a4.split(": ")[-1]
              and len(a4.split(": ")[-1].split(",")[-1]) == 4, a4)
        a5 = str(ws["A5"].value)
        pf = m2["proxima_fecha_cierre"]
        check("A5 dice índice del próximo cierre",
              pf[8:10] + "/" + pf[5:7] + "/" + pf[:4] in a5, a5)
        check("B6 = coeficiente", ws["B6"].value is not None
              and abs(float(ws["B6"].value) - m2["coeficiente"]) < 1e-9,
              str(ws["B6"].value))
        check("B6 formateado4 decimales", "0.0000" in str(ws["B6"].number_format),
              str(ws["B6"].number_format))
        encabezados = [str(c.value) for c in ws[9]]
        check("encabezado con Saldo del cierre",
              any(fecha_fin[8:10] + "/" + fecha_fin[5:7] + "/" + fecha_fin[:4] in e
                  for e in encabezados), str(encabezados))
        check("encabezado con Actualizado del próximo cierre",
              any(pf[8:10] + "/" + pf[5:7] + "/" + pf[:4] in e for e in encabezados),
              str(encabezados))

        st, _ = req(
            "PUT",
            "/api/balance-rt54/ejercicios/1",
            {
                "cabecera": ej["cabecera"],
                "valores": [{
                    "seccion": "esp",
                    "clave": fila_rubro["clave"],
                    "valor_actual": original_actual,
                    "valor_anterior": original_anterior,
                }],
                "celdas_nota": [],
            },
            token,
        )
        check("restaura importe original", st == 200, str(st))

# ------------------------------------- ancla del cliente (otro cliente) -----
# Se usa un cliente cuya ficha se pueda guardar (el13 es RI sin alícuota).
st, clientes = req("GET", "/api/clientes", token=token)
cand = next(
    (c for c in clientes
     if c["condicion_iva"] != "responsable_inscripto" or c.get("alicuota_iva")),
    None,
)
check("hay cliente apto para la ancla", cand is not None, str(len(clientes or [])))
if cand:
    st, cli = req("GET", f"/api/clientes/{cand['id']}", token=token)
    p = payload_cliente(cli)
    ancla_original = p["fecha_cierre_ejercicio"]

    p["fecha_cierre_ejercicio"] = "2026-08-31"
    st, g = req("PUT", f"/api/clientes/{cand['id']}", p, token)
    check("guarda fecha de cierre en el cliente", st == 200
          and g.get("fecha_cierre_ejercicio") == "2026-08-31",
          f"{st} {str(g)[:200]}")

    # Ejercicio que cierra30/11/2025: con ancla31/08 el próximo cierre es
    # 31/08/2026 (y NO30/11/2026, que daría el cálculo automático).
    st, nuevo = req("POST", "/api/balance-rt54/ejercicios", {
        "cliente_id": cand["id"],
        "nombre": "prueba-ancla",
        "fecha_inicio": "2025-01-01",
        "fecha_fin": "2025-11-30",
    }, token)
    check("crea ejercicio de prueba", st in (200, 201), str(st) + str(nuevo)[:200])
    if st in (200, 201):
        ej_id = nuevo["cabecera"]["id"]
        st, ma = req("GET", f"/api/balance-rt54/ejercicios/{ej_id}/moneda-homogenea",
                     token=token)
        check("moneda con ancla responde", st == 200, str(st) + str(ma)[:200])
        if st == 200:
            check("próximo cierre según ancla del cliente (31/08/2026)",
                  ma["proxima_fecha_cierre"] == "2026-08-31",
                  ma["proxima_fecha_cierre"])
            check("índices de nov-2025 y ago-2026",
                  abs(ma["indice_anterior"]["valor"]
                      - float(cuatro(por_mes["2025-11-01"]["indice"]))) < 1e-9
                  and abs(ma["indice_nuevo"]["valor"]
                          - float(cuatro(por_mes["2026-08-01"]["indice"]))) < 1e-9,
                  f"{ma['indice_anterior']} {ma['indice_nuevo']}")

        # ancla a un mes sin índice cargado = aviso claro
        p["fecha_cierre_ejercicio"] = "2026-09-30"
        st, _ = req("PUT", f"/api/clientes/{cand['id']}", p, token)
        check("cambia ancla a sep-2026", st == 200, str(st))
        st, r = req("GET", f"/api/balance-rt54/ejercicios/{ej_id}/moneda-homogenea",
                    token=token)
        check("falta índice de sep-2026 y avisa", st == 400
              and "septiembre de 2026" in str(r.get("detail", r)),
              f"{st} {str(r)[:200]}")

        st, _ = req("DELETE", f"/api/balance-rt54/ejercicios/{ej_id}", token=token)
        check("borra ejercicio de prueba", st == 204, str(st))

    p["fecha_cierre_ejercicio"] = ancla_original
    st, _ = req("PUT", f"/api/clientes/{cand['id']}", p, token)
    check("restaura fecha de cierre original", st == 200, str(st))

print(f"\n{ok} OK / {len(fallos)} fallos")
if fallos:
    raise SystemExit(1)
