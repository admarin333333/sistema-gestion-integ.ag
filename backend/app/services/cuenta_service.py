from datetime import date
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Font
from sqlalchemy import Integer, Numeric, String, func, literal, select, union_all, literal_column
from sqlalchemy.orm import Session

from app.config import settings
from app.models.anticipo import Anticipo
from app.models.factura import Factura
from app.models.recibo import Recibo
from app.schemas.factura import ETIQUETAS_TIPO
from app.schemas.recibo import CuentaCorrienteOut, MovimientoOut
from app.services import cliente_service

CERO = literal(0, Numeric(14, 2))
# orden de los movimientos de un mismo día: primero lo que se DEBE,
# después lo que se pagó (HABER) y al final las anulaciones.
DEBE_PRIMERO = literal(0, Integer)
HABER_DESPUES = literal(1, Integer)
ANULACION = literal(2, Integer)


def _movimientos(db: Session, cliente_id: int):
    """Facturas (DEBE) + recibos y anticipos (HABER) del cliente, en una tabla."""
    facturas = (
        select(
            Factura.fecha.label("fecha"),
            literal("factura", String(10)).label("origen"),
            DEBE_PRIMERO.label("orden"),
            Factura.id.label("ref_id"),
            Factura.tipo_comprobante.label("tipo"),
            func.concat(Factura.punto_venta, "-", Factura.numero).label("numero"),
            Factura.importe.label("debe"),
            CERO.label("haber"),
        ).where(Factura.cliente_id == cliente_id, Factura.estado != "anulada")
    )
    recibos = select(
        Recibo.fecha,
        literal("recibo", String(10)),
        HABER_DESPUES,
        Recibo.id,
        literal("", String(10)),
        Recibo.numero,
        CERO,
        Recibo.importe,
    ).where(Recibo.cliente_id == cliente_id)
    anticipos = select(
        Anticipo.fecha,
        literal("anticipo", String(10)),
        HABER_DESPUES,
        Anticipo.id,
        literal("", String(10)),
        Anticipo.numero,
        CERO,
        Anticipo.importe,
    ).where(Anticipo.cliente_id == cliente_id)
    # Un anticipo eliminado genera su movimiento contrario en DEBE, para que
    # el estado de cuenta muestre qué pasó (la plata se devolvió).
    anticipos_eliminados = select(
        Anticipo.fecha,
        literal("anticipo_eliminado", String(24)),
        ANULACION,
        Anticipo.id,
        literal("", String(10)),
        Anticipo.numero,
        Anticipo.importe,
        CERO,
    ).where(
        Anticipo.cliente_id == cliente_id,
        Anticipo.estado == "eliminado",
    )
    return union_all(
        facturas, recibos, anticipos, anticipos_eliminados
    ).subquery("mov")


def _condiciones_fecha(m, desde, hasta) -> list:
    """Condiciones de fecha usando la columna del subquery (m.c.fecha)."""
    condiciones = []
    if desde:
        condiciones.append(m.c.fecha >= desde)
    if hasta:
        condiciones.append(m.c.fecha <= hasta)
    return condiciones


def movimientos(
    db: Session,
    cliente_id: int,
    desde=None,
    hasta=None,
) -> CuentaCorrienteOut:
    """Historial con saldo corrido y totales — todo calculado con SQL.

    Imputar un anticipo NO agrega un movimiento nuevo: la plata ya entró
    como HABER al alta, la imputación solo marca a qué factura va.
    """
    cliente = cliente_service.obtener(db, cliente_id)
    m = _movimientos(db, cliente_id)

    # Aplicar filtros de fecha sobre la subquery
    condiciones_fecha = _condiciones_fecha(m, desde, hasta)
    if condiciones_fecha:
        # m es Subquery -> usar select(m).where(...) para filtrar
        m = select(m).where(*condiciones_fecha).subquery("mov_filtrado")

    consulta = select(
        m.c.fecha,
        m.c.origen,
        m.c.tipo,
        m.c.numero,
        m.c.debe,
        m.c.haber,
        # saldo corrido: SUM(DEBE − HABER) hasta esa fila, en la base de datos
        func.sum(m.c.debe - m.c.haber)
        .over(order_by=[m.c.fecha, m.c.orden, m.c.ref_id])
        .label("saldo"),
    ).order_by(m.c.fecha, m.c.orden, m.c.ref_id)

    filas = []
    for fila in db.execute(consulta):
        if fila.origen == "factura":
            concepto = f"{ETIQUETAS_TIPO.get(fila.tipo, fila.tipo)} {fila.numero}"
        elif fila.origen == "anticipo_eliminado":
            concepto = f"Eliminación anticipos {fila.numero}"
        elif fila.origen == "anticipo":
            concepto = f"Anticipo {fila.numero}"
        else:
            concepto = f"Recibo {fila.numero}"
        filas.append(
            MovimientoOut(
                fecha=fila.fecha,
                concepto=concepto,
                debe=float(fila.debe or 0),
                haber=float(fila.haber or 0),
                saldo=float(fila.saldo or 0),
            )
        )

    total_debe, total_haber = db.execute(
        select(
            func.coalesce(func.sum(m.c.debe), 0),
            func.coalesce(func.sum(m.c.haber), 0),
        )
    ).one()

    debe = float(total_debe)
    haber = float(total_haber)
    return CuentaCorrienteOut(
        cliente_id=cliente.id,
        cliente_nombre=cliente.nombre_completo,
        movimientos=filas,
        total_debe=debe,
        total_haber=haber,
        saldo=debe - haber,
    )


# --------------------------------------------------------------------- Excel

def informe_excel(
    cliente,
    movimientos: list[MovimientoOut],
    total_debe: float,
    total_haber: float,
    saldo: float,
    desde=None,
    hasta=None,
) -> bytes:
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Estado de cuenta"

    # Cabecera del cliente
    hoja.append([settings.nombre_estudio])
    hoja.cell(row=1, column=1).font = Font(bold=True, size=13)
    hoja.append(["ESTADO DE CUENTA"])
    hoja.cell(row=2, column=1).font = Font(bold=True, size=11)
    hoja.append([])

    # Datos del cliente
    hoja.append(["Cliente:", cliente.nombre_completo])
    if cliente.tipo_persona == "fisica":
        hoja.append(["CUIT:", cliente.cuit or "—"])
        hoja.append(["DNI:", cliente.dni or "—"])
    else:
        hoja.append(["CUIT:", cliente.cuit or "—"])
    hoja.append(["Condición IVA:", cliente.tipo_actividad.replace("_", " ").title()])
    hoja.append(["Email:", cliente.email or "—"])
    tel = " ".join(filter(None, [cliente.cod_area, cliente.telefono]))
    hoja.append(["Teléfono:", tel or "—"])
    hoja.append([])

    # Rango de fechas
    if desde or hasta:
        rango = f"Desde: {desde or 'inicio'}  —  Hasta: {hasta or 'hoy'}"
        hoja.append([rango])
        hoja.append([])

    # Tabla de movimientos
    columnas = [
        ("Fecha", 12),
        ("Concepto", 50),
        ("DEBE", 14),
        ("HABER", 14),
        ("SALDO", 14),
    ]
    hoja.append([nombre for nombre, _ in columnas])
    for celda in hoja[hoja.max_row]:
        celda.font = Font(bold=True)

    for mov in movimientos:
        hoja.append(
            [
                mov.fecha.isoformat() if mov.fecha else "",
                mov.concepto,
                mov.debe,
                mov.haber,
                mov.saldo,
            ]
        )

    # Totales
    hoja.append([""] * len(columnas))
    hoja.append(["", "TOTALES", total_debe, total_haber, saldo])
    ultima = hoja.max_row
    for col in range(2, 6):
        hoja.cell(row=ultima, column=col).font = Font(bold=True)
        hoja.cell(row=ultima, column=col).number_format = "#,##0.00"

    for (nombre, ancho), letra in zip(columnas, "ABCDE"):
        hoja.column_dimensions[letra].width = ancho

    salida = BytesIO()
    libro.save(salida)
    return salida.getvalue()
