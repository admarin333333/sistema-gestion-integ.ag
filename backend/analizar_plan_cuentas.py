"""Analiza el Excel del plan de cuentas y reporta qué se puede cargar y qué falta.

No escribe nada en la base: solo informa, para decidir con el usuario.

Uso:  python -X utf8 analizar_plan_cuentas.py
"""

import re

import openpyxl

ARCHIVO = r"C:\estudio contable\plan de cuentas.xlsx"

# "1.1.01.01 Caja"  ->  código "1.1.01.01", nombre "Caja"
RE_CODIGO = re.compile(r"^(\d+(?:[.,]\d+)*)[.\s]+(.+)$")
RE_SOLO_NUM = re.compile(r"^(\d+(?:[.,]\d+)*)[.\s]*$")


def limpiar(v):
    return "" if v is None else str(v).strip()


def analizar(hoja):
    filas = []
    for i, row in enumerate(hoja.iter_rows(values_only=True), 1):
        codigo_crudo, nombre, imputable = limpiar(row[0]), limpiar(row[1]), limpiar(row[2])
        if not nombre or not (codigo_crudo or nombre):
            continue
        # El código a veces viene en la columna 1 y a veces dentro del nombre.
        codigo = codigo_crudo.replace(",", ".")
        nombre_limpio = nombre
        m = RE_CODIGO.match(nombre)
        if m:
            codigo = m.group(1).replace(",", ".")
            nombre_limpio = m.group(2).strip()
        elif RE_SOLO_NUM.match(nombre):
            codigo = RE_SOLO_NUM.match(nombre).group(1).replace(",", ".")
            nombre_limpio = ""
        filas.append(
            {
                "fila": i,
                "codigo": codigo,
                "nombre": nombre_limpio,
                "imputable": imputable.lower() in ("si", "sí"),
                "tiene_col_codigo": bool(codigo_crudo),
            }
        )
    return filas


def nivel(codigo):
    return len(codigo.split(".")) if codigo else 0


def main():
    wb = openpyxl.load_workbook(ARCHIVO, data_only=True)

    for hoja in wb.sheetnames:
        filas = analizar(wb[hoja])
        sin_codigo = [f for f in filas if not f["codigo"]]
        sin_nombre = [f for f in filas if not f["nombre"]]
        imputables = [f for f in filas if f["imputable"]]
        # Cuenta con hijos: el padre no debería ser imputable.
        codigos = {f["codigo"] for f in filas if f["codigo"]}
        # Una cuenta "tiene hijos" si OTRO código empieza con el suyo + punto.
        padres_imputables = [
            f
            for f in filas
            if f["codigo"]
            and f["imputable"]
            and any(
                otro != f["codigo"] and otro.startswith(f["codigo"] + ".")
                for otro in codigos
            )
        ]

        print("=" * 74)
        print(f"HOJA '{hoja}': {len(filas)} filas con contenido")
        print("=" * 74)
        print(f"  Con código ........... {len(filas) - len(sin_codigo)}")
        print(f"  SIN código (no se dónde ubicarlos): {len(sin_codigo)}")
        print(f"  Con nombre en blanco .. {len(sin_nombre)}")
        print(f"  Marcadas imputable=sí . {len(imputables)}")
        print(f"  Imputables que TIENEN hijos (no deberían): {len(padres_imputables)}")
        print()
        if sin_codigo:
            print("  Filas sin código (necesitan código nuevo):")
            for f in sin_codigo:
                print(f"    fila {f['fila']:>3} | {f['nombre']}")
        print()
        if padres_imputables:
            print("  Imputables que ya tienen subcuentas:")
            for f in padres_imputables[:12]:
                print(f"    {f['codigo']} {f['nombre']}")
            if len(padres_imputables) > 12:
                print(f"    ... y {len(padres_imputables) - 12} más")
        print()


if __name__ == "__main__":
    main()